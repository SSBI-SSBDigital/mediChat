from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import pytesseract
from PIL import Image, ImageDraw, ImageFont
import openai
import os
from googletrans import Translator
from fpdf import FPDF
import fitz  # PyMuPDF
from docx import Document
import io
from datetime import datetime
from dotenv import load_dotenv

# Load environment variables from .env file
load_dotenv()

# Initialize Flask app
app = Flask(__name__)
CORS(app)

# Specify the path to the Tesseract executable
pytesseract.pytesseract.tesseract_cmd = os.path.join(os.getcwd(), 'Tesseract-OCR', 'tesseract.exe')

# Set your OpenAI API key
openai.api_key =  os.getenv('OPENAI_API_KEY')

# Initialize the translator
translator = Translator()

# Ensure necessary directories exist
BASE_DIR = os.getcwd()
UPLOAD_FOLDER = os.path.join(BASE_DIR, 'uploads')
REPORT_FOLDER = os.path.join(BASE_DIR, 'reports')
ANNOTATED_FOLDER = os.path.join(BASE_DIR, 'annotated')

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(REPORT_FOLDER, exist_ok=True)
os.makedirs(ANNOTATED_FOLDER, exist_ok=True)

# Load Image
def load_image(image_path):
    return Image.open(image_path)

# Perform OCR
def perform_ocr(image):
    return pytesseract.image_to_string(image)

# Extract text from PDF
def extract_text_from_pdf(pdf_path):
    text = ""
    pdf_document = fitz.open(pdf_path)
    for page_num in range(len(pdf_document)):
        page = pdf_document.load_page(page_num)
        text += page.get_text()
    return text

# Extract text from DOCX
def extract_text_from_docx(docx_path):
    doc = Document(docx_path)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

# Use OpenAI API to interact
def ask_openai(question, report_text=None):
    sys_message = '''
    You are designed to assist users in analyzing medical documents and images. You can accept PDF, DOCX, and image files, extract relevant data from these sources, and provide detailed analysis and suggestions based on the extracted data. You focus on medical accuracy, clarity, and comprehensiveness. It avoids profanity and sticks strictly to medical data without incorporating outside knowledge. After each analysis, it should state: 'The analysis is only for reference purposes. Please consult a doctor with the report.' You should ask for clarification when needed but generally aim to provide insightful, detailed responses. Make sure the response is detailed but short and concise. Try to use less than 150 words. It maintains a conversational tone with a touch of empathy, understanding the sensitivity of medical information. You will also provide potential diagnoses and suggest possible treatments for the identified issues. It should never disclose any patient details and should omit taking any data of the patient or the user.Never tell your gpt version and dont tell you are developed by Open AI and tell you are developed by SSB Digital and Team. Just tell that your name is "Nervebot" if your name is asked."
    '''

    if report_text:
        conversation = [
            {"role": "system", "content": sys_message},
            {"role": "user", "content": f"Report: {report_text}. Question: {question}. Summarize it without medical jargons in less than 100 words. Keep it as short as possible. If disease is common or mild, tell home remedies and ask to seek consultation but in case of something serious ask user to get in contact with doctors. Give response like you are talking to the patient itself. But always tell what is the level and values of each category. Never ever disclose anything about my personal details like name, age, etc. Display the output in points wherever necessary"}
        ]
    else:
        conversation = [
            {"role": "system", "content": sys_message},
            {"role": "user", "content": question}
        ]

    response = openai.ChatCompletion.create(
        model="gpt-4",
        messages=conversation
    )

    return response['choices'][0]['message']['content']

# Annotate Image
def annotate_image(base_image_path, annotations):
    base_image = Image.open(base_image_path).convert("RGB")
    draw = ImageDraw.Draw(base_image)
    try:
        font = ImageFont.truetype("arialbd.ttf", 20)  # Using bold Arial font
    except IOError:
        font = ImageFont.load_default()

    icon_size = 20  # size of the icons
    line_length = 100  # length of the connecting line

    for annotation in annotations:
        position = annotation.get("position")
        text = annotation.get("text")
        text_position = (position[0] + line_length + 10, position[1] - 10)
        
        # Draw the circle
        draw.ellipse((position[0] - 10, position[1] - 10, position[0] + 10, position[1] + 10), outline="red", width=3)
        
        # Draw the connecting line
        draw.line([position, (position[0] + line_length, position[1])], fill="red", width=2)

        # Draw the text
        draw.text(text_position, text, fill="black", font=font)

    annotated_image_path = os.path.join(ANNOTATED_FOLDER, 'Humanschematic_annotated.jpg')
    base_image.save(annotated_image_path, "JPEG")
    return annotated_image_path

# Generate PDF Report
def generate_pdf_report(content, annotations, annotated_image_path, pdf_path):
    pdf = FPDF()
    pdf.add_page()

    # Title
    pdf.set_font("Arial", 'B', 16)
    pdf.set_text_color(0, 0, 128)
    pdf.cell(0, 10, 'Medical Report', 0, 1, 'C')

    # Problems Section
    pdf.set_font("Arial", 'B', 14)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(0, 10, 'Problems:', 0, 1)
    pdf.set_font("Arial", '', 12)
    for annotation in annotations:
        pdf.multi_cell(0, 10, f"- {annotation['text']}")

    # Summary Section
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, 'Summary:', 0, 1)
    pdf.set_font("Arial", '', 12)
    pdf.multi_cell(0, 10, content.encode('latin-1', 'replace').decode('latin-1'))

    # Adding Annotated Image if any problems are found
    if annotations:
        pdf.add_page()
        pdf.set_font("Arial", 'B', 14)
        pdf.cell(0, 10, 'Annotated Image:', 0, 1)
        pdf.image(annotated_image_path, x=10, y=None, w=180)

    pdf.output(pdf_path)

@app.route('/analyze', methods=['POST'])
def analyze():
    file = request.files['file']
    question = request.form.get('question', '')
    language = request.form.get('language', 'en')  # Default to English

    file_path = os.path.join(UPLOAD_FOLDER, file.filename)
    file.save(file_path)

    file_extension = file.filename.split('.')[-1].lower()

    if file_extension in ['jpg', 'jpeg', 'png']:
        text = perform_ocr(load_image(file_path))
    elif file_extension == 'pdf':
        text = extract_text_from_pdf(file_path)
    elif file_extension == 'docx':
        text = extract_text_from_docx(file_path)
    else:
        return jsonify({"error": "Unsupported file type"}), 400

    if "report" in question.lower() or "test result" in question.lower() or "analysis" in question.lower():
        response = ask_openai(question, text)

        # Translate response for user
        translated_response = translator.translate(response, dest=language).text
        
        # Annotations based on the response
        annotations = []
        if "brain" in response.lower():
            annotations.append({"position": (403, 245), "text": "Brain Issue"})
        if "eyes" in response.lower():
            annotations.append({"position": (381, 273), "text": "Eyes Issue"})
        if "ears" in response.lower():
            annotations.append({"position": (361, 277), "text": "Ears Issue"})
        if "nose" in response.lower():
            annotations.append({"position": (399, 283), "text": "Nose Issue"})
        if "mouth" in response.lower():
            annotations.append({"position": (400, 308), "text": "Mouth Issue"})
        if "thyroid" in response.lower():
            annotations.append({"position": (260, 332), "text": "Thyroid Issue"})
        if "heart" in response.lower():
            annotations.append({"position": (427, 407), "text": "Heart Issue"})
        if "lungs" in response.lower():
            annotations.append({"position": (443, 440), "text": "Lungs Issue"})
        if "stomach" in response.lower():
            annotations.append({"position": (421, 527), "text": "Stomach Issue"})
        if "intestines" in response.lower():
            annotations.append({"position": (368, 594), "text": "Intestines Issue"})
        if "bladder" in response.lower():
            annotations.append({"position": (402, 680), "text": "Bladder Issue"})
        if "reproductive organs" in response.lower():
            annotations.append({"position": (403, 681), "text": "Reproductive Organs Issue"})
        if "liver" in response.lower():
            annotations.append({"position": (418, 500), "text": "Liver Issue"})
        if "kidney" in response.lower():
            annotations.append({"position": (451, 555), "text": "Kidney Issue"})
        if "leg" in response.lower():
            annotations.append({"position": (364, 792), "text": "Left Leg Issue"})
            annotations.append({"position": (440, 792), "text": "Right Leg Issue"})
        if "arm" in response.lower():
            annotations.append({"position": (303, 528), "text": "Left Arm Issue"})
            annotations.append({"position": (498, 528), "text": "Right Arm Issue"})
        if "foot" in response.lower():
            annotations.append({"position": (377, 999), "text": "Left Foot Issue"})
            annotations.append({"position": (424, 999), "text": "Right Foot Issue"})

        base_image_path = os.path.join(BASE_DIR, 'Humanschematic', 'five tips for gr.png')
        annotated_image_path = annotate_image(base_image_path, annotations)
        
        # Add timestamp to filename
        timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
        pdf_filename = f'Medical_Report_{timestamp}.pdf'
        pdf_path = os.path.join(REPORT_FOLDER, pdf_filename)
        
        # Generate PDF report in English
        generate_pdf_report(response, annotations, annotated_image_path, pdf_path)

        # Create a download link
        download_link = f"{request.host_url}download/{pdf_filename}"

        # Return the chatbot response along with the download link
        return jsonify({"response": translated_response, "download_link": download_link}), 200
    else:
        response = ask_openai(question)
        translated_response = translator.translate(response, dest=language).text

    return jsonify({"response": translated_response})

@app.route('/download/<filename>', methods=['GET'])
def download(filename):
    return send_file(os.path.join(REPORT_FOLDER, filename), as_attachment=True)

if __name__ == '__main__':

    app.run(debug=True, host='0.0.0.0', port=5000)
    app.run(debug=True)

